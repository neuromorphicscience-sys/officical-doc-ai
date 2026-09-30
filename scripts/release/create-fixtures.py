"""Create synthetic, public release samples. Requires python-docx and Pillow."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
from io import BytesIO
import json, zipfile, re
ROOT=Path(__file__).resolve().parents[2]
DEMO=ROOT/'demo'; CASES=DEMO/'cases'; CASES.mkdir(parents=True,exist_ok=True)
notice=['关于开展办公文档规范化工作的通知','各学院、各有关单位：','为提高日常办公文档质量，减少重复排版工作，现就开展文档规范化工作通知如下。','一、总体要求','（一）统一工作原则','坚持内容准确、结构清晰、格式规范。文档整理应保留原有文字，不擅自增加、删除或改写正文。','（二）明确责任分工','各单位指定联络人员，负责收集样本文档、核对结构识别结果，并记录需要人工确认的问题。','二、主要任务','（一）开展样本整理','选取通知、报告等常见文种，重点检查标题层级、正文段落、附件说明和落款日期。','（二）做好结果复核','规范化完成后，核对字符数量和完整性校验结果，使用办公软件打开输出文档并检查实际分页。','三、组织保障','请各单位于十月底前完成首轮整理，形成可复用的示例材料。本通知仅用于软件演示，不涉及真实业务安排。','附件：办公文档检查事项清单','示例大学办公室','2026年9月30日']
unnum=['关于推进办公资料整理工作的通知','各有关单位：','为便于资料查阅和工作交接，现对办公资料整理提出以下要求。','总体要求','坚持真实完整、分类清晰。保留原始文档内容，只整理文档结构与版式。','主要任务','统一整理归档资料，检查标题与正文逻辑，逐项核对附件、落款和日期。','组织保障','建立定期复核机制，由专人检查整理结果并反馈异常情况。','示例大学办公室','2026年9月30日']
attach=['关于报送示例材料的通知','各有关单位：','请按要求报送示例材料，并确保所列信息完整准确。本材料仅用于软件测试。','附件：示例材料目录','示例大学办公室','2026年9月30日','（此件公开发布）']
levels=['关于规范资料报送流程的通知','各有关单位：','请按以下层级检查各项任务。','一、工作安排','（一）材料准备','1. 核对清单','（1）检查附件','各项资料应保持真实、完整。','示例大学办公室','2026年9月30日']
long=['关于开展年度办公资料整理的报告','一、工作情况']+[f'第{i+1}批材料已完成核对。各单位按照统一要求整理资料，复查文件内容、责任分工与归档信息，确保文字完整、流程清晰。' for i in range(65)]+['二、下一步工作','继续完善材料复核机制，提高文档整理效率。','示例大学办公室','2026年9月30日']
letter=['关于商洽示例交流活动的函','示例研究中心：','为交流文档整理经验，拟组织一次示例研讨活动，现函商有关事项。','一、交流安排']+[f'交流事项{i+1}：请协助确认会务安排和资料准备工作，本段为用于分页验收的虚构测试文字，不涉及任何真实单位业务。' for i in range(22)]+['二、有关要求','请予以支持，并函复有关安排。','示例大学办公室','2026年9月30日']
amb=['工作情况说明','各有关单位：','现将近期工作情况说明如下。','再作说明','这一短语可能是正文的承接语，也可能是新的小节标题，需要结合具体上下文人工核对。','其他','后续事项待讨论。本测试用于展示不确定性提示，不预设模型必须产生某个置信度。']
manifest=[]
def create(name,paras,messy=False,table=False,image=False,header=False):
 d=Document();d.core_properties.author='Release QA';d.core_properties.last_modified_by='Release QA';d.core_properties.title='合成演示文档'
 for i,t in enumerate(paras):
  p=d.add_paragraph();p.style=d.styles['Normal'];r=p.add_run(t);r.font.name='Microsoft YaHei';r.font.size=Pt([11,18,9,14,12][i%5] if messy else 12)
  r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
  if messy:
   r.font.color.rgb=RGBColor.from_string(['394A6D','A34238','225E89','596577'][i%4]);r.bold=i%3==0
   p.paragraph_format.left_indent=Pt((i%3)*8);p.paragraph_format.space_after=Pt((i%4)*5)
   p.alignment=[0,2,1,0][i%4]
  else: p.paragraph_format.space_after=Pt(0)
 if table:
  t=d.add_table(rows=3,cols=2);t.style='Table Grid'
  for row,vals in zip(t.rows,[['检查事项','状态'],['正文完整性','待核对'],['附件与日期','待核对']]):
   for cell,v in zip(row.cells,vals):cell.text=v
 if image:
  im=Image.new('RGB',(640,200),'#edf3fb');draw=ImageDraw.Draw(im);draw.rectangle((28,35,612,165),outline='#2458a6',width=4);draw.line((60,100,580,100),fill='#2458a6',width=5)
  buf=BytesIO();im.save(buf,format='PNG');buf.seek(0);d.add_picture(buf,width=Inches(4.8))
 if header:
  d.sections[0].header.paragraphs[0].text='合成文档 · 页眉保留测试'
  d.sections[0].footer.paragraphs[0].text='合成文档 · 页脚保留测试'
 path=CASES/name;d.save(path)
 # happy-dom's XML declaration parser requires double quotes; both forms are
 # legal XML. Keep the real browser demo in python-docx's original serialization.
 if name != 'demo-temp.docx':
  with zipfile.ZipFile(path) as z: parts={n:z.read(n) for n in z.namelist()}
  with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
   for n,data in parts.items():
    if n.endswith(('.xml','.rels')): data=re.sub(br'<\?xml[^?]*\?>',lambda m:m.group().replace(b"'",b'"'),data,count=1)
    z.writestr(n,data)
 return path
spec=[('01-uniform.docx',notice,{}),('02-numbering.docx',levels,{'messy':True}),('03-unnumbered.docx',unnum,{}),('04-attachments.docx',attach,{}),('05-long.docx',long,{'header':True}),('06-letter.docx',letter,{'header':True}),('07-table.docx',attach,{'table':True,'header':True}),('08-image.docx',attach,{'image':True,'header':True}),('09-ambiguous.docx',amb,{})]
for n,ps,kw in spec:
 create(n,ps,**kw);manifest.append({'file':n,'paragraphs':len(ps),'kind':n.split('-')[1].split('.')[0]})
(CASES/'10-corrupt.docx').write_bytes(b'Intentional invalid ZIP for error recovery test.');manifest.append({'file':'10-corrupt.docx','expected':'reject'})
p=create('demo-temp.docx',notice,messy=True);p.replace(DEMO/'乱格式办公通知示例.docx')
(CASES/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
print('Created 10 adversarial cases and demo DOCX; all content synthetic.')
